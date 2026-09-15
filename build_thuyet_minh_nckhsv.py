from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
from zipfile import ZipFile

OUT = Path("BATS-AgriGuard_Du_thao_Thuyet_minh_NCKHSV_2026.docx")
EUREKA_SOURCE = Path("/Users/nguyendinhkhang/Downloads/Hồ sơ đăng ký Eureka 2026/Thuyết minh Eureka 2026.docx")
ASSET_DIR = Path("/private/tmp/bats_eureka_assets")

BLUE = "0000FF"
GRAY = "F2F4F7"
LIGHT_BLUE = "D9EAF7"

def set_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_width(cell, width_dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.first_child_found_in("w:tblInd")
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "0")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def no_cell_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def add_text(p, text, **kwargs):
    r = p.add_run(text)
    set_font(r, **kwargs)
    return r

def add_para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, indent=0, first_line=0, size=12, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.12
    pf.left_indent = Inches(indent)
    pf.first_line_indent = Inches(first_line)
    if text:
        add_text(p, text, size=size, bold=bold, italic=italic, color=color)
    return p

def add_heading(doc, text, level=1):
    sizes = {1: 13.5, 2: 12.5, 3: 11.5}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(4)
    pf.keep_with_next = True
    add_text(p, text, size=sizes[level], bold=True, color=BLUE)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    add_text(p, text, size=11)
    return p

def fill_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=11):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    add_text(p, text, size=size, bold=bold)

def add_kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], GRAY)
        fill_cell(cells[0], label, bold=True, size=11)
        fill_cell(cells[1], value, size=11)
    set_table_geometry(table, [2550, 6810])
    return table

def add_grid_table(doc, headers, rows, widths, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE)
        fill_cell(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            fill_cell(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 30 else WD_ALIGN_PARAGRAPH.LEFT, size=font_size)
    set_table_geometry(table, widths)
    return table

def add_reference(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent = Inches(0.32)
    pf.first_line_indent = Inches(-0.32)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    add_text(p, f"[{number}] {text}", size=10.5)

def prepare_eureka_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    wanted = ("image1.png",)
    with ZipFile(EUREKA_SOURCE) as archive:
        for filename in wanted:
            target = ASSET_DIR / filename
            target.write_bytes(archive.read(f"word/media/{filename}"))

def add_figure(doc, filename, caption, width=6.15):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(ASSET_DIR / filename), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    add_text(cap, caption, size=10.5, italic=True)

doc = Document()
prepare_eureka_assets()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(11)

# Title block follows the supplied student-research proposal structure.
top = doc.add_table(rows=1, cols=2)
top.alignment = WD_TABLE_ALIGNMENT.CENTER
no_cell_borders(top)
fill_cell(top.cell(0, 0), "ĐẠI HỌC QUỐC GIA TP. HỒ CHÍ MINH\nTRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
fill_cell(top.cell(0, 1), "Ngày nhận hồ sơ:\n\n(Do CQ quản lý ghi)", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
set_table_geometry(top, [5200, 4160])

add_para(doc, "THUYẾT MINH", align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=1, size=16, bold=True)
add_para(doc, "ĐỀ TÀI NGHIÊN CỨU KHOA HỌC SINH VIÊN 2026", align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=14, size=14, bold=False)

add_heading(doc, "A. THÔNG TIN CHUNG", 1)
add_heading(doc, "A1. Tên đề tài", 2)
add_para(doc, "- Tên tiếng Việt (IN HOA): BATS-AGRIGUARD: HỆ THỐNG PHÁT HIỆN VÀ CẢNH BÁO NGUY CƠ LẠM DỤNG MÃ SỐ VÙNG TRỒNG TRONG TRUY XUẤT NGUỒN GỐC CHUỖI NÔNG SẢN", after=4, first_line=0, size=11)
add_para(doc, "- Tên tiếng Anh (IN HOA): BATS-AGRIGUARD: A RISK-AWARE FRAMEWORK FOR DETECTING AND PREVENTING PLANTATION CODE MISUSE IN AGRICULTURAL SUPPLY CHAIN TRACEABILITY", after=6, first_line=0, size=11)

add_heading(doc, "A2. Thời gian thực hiện", 2)
add_para(doc, "6 tháng (kể từ khi được duyệt).", size=11)

add_heading(doc, "A3. Tổng kinh phí", 2)
add_para(doc, "Tổng kinh phí: 1.000.000 đồng, gồm:", after=3, size=11)
add_bullet(doc, "Kinh phí từ Trường Đại học Công nghệ Thông tin: 1.000.000 đồng.")
add_para(doc, "Tài liệu nguồn chưa cung cấp phương án phân bổ chi tiết theo hạng mục.", after=5, size=10.5, italic=True)

add_heading(doc, "A4. Chủ nhiệm đề tài", 2)
add_para(doc, "Họ và tên: Nguyễn Đình Khang", after=2, indent=0.22, size=11)
add_para(doc, "Ngày, tháng, năm sinh: 07/05/2005                                      Giới tính (Nam/Nữ): Nam", after=2, indent=0.22, size=11)
add_para(doc, "Số CCCD: 080205003038; Ngày cấp: 25/04/2021; Nơi cấp: Cục trưởng Cục Cảnh sát quản lý hành chính về trật tự xã hội", after=2, indent=0.22, size=11)
add_para(doc, "Mã số sinh viên: 23520694", after=2, indent=0.22, size=11)
add_para(doc, "Số điện thoại liên lạc: 0967 414 607", after=2, indent=0.22, size=11)
add_para(doc, "Khoa: Mạng máy tính và Truyền thông", after=2, indent=0.22, size=11)
add_para(doc, "Số tài khoản: 1068762359                                      Ngân hàng: Vietcombank", after=2, indent=0.22, size=11)
add_para(doc, "Giảng viên hướng dẫn: [Chưa có thông tin trong tài liệu nguồn]", after=5, indent=0.22, size=10.5)

doc.add_page_break()
add_heading(doc, "A5. Thành viên đề tài (kể cả chủ nhiệm)", 2)
add_grid_table(doc, ["TT", "Họ và tên", "MSSV", "Khoa"], [
    ["1", "Nguyễn Đình Khang", "23520694", "Mạng máy tính và Truyền thông"],
    ["2", "Võ Minh An", "23520033", "Mạng máy tính và Truyền thông"],
], [650, 2850, 1800, 4060], 10.5)

add_heading(doc, "B. MÔ TẢ NGHIÊN CỨU", 1)
add_heading(doc, "B1. Giới thiệu về đề tài", 2)
add_para(doc, "Truy xuất nguồn gốc nông sản đang được nghiên cứu theo hướng kết hợp blockchain, Internet of Things và mô hình dữ liệu chuẩn hóa để tăng khả năng chia sẻ, kiểm toán và minh bạch lịch sử sản phẩm trong chuỗi cung ứng [1]-[4]. Các công trình hiện có cho thấy blockchain có thể bảo toàn lịch sử giao dịch và hỗ trợ phối hợp giữa nhiều tác nhân; những triển khai thực tế cũng đã chứng minh tiềm năng áp dụng trong chuỗi nông sản [3], [4].")
add_para(doc, "Trong một chuỗi nông sản có nhiều tác nhân, dữ liệu vùng trồng, thu hoạch, tiếp nhận, chuyển giao, đóng gói và xuất hàng thường được hình thành ở các thời điểm và hệ thống khác nhau. Vì vậy, yêu cầu không chỉ là lưu lại các bản ghi riêng lẻ, mà còn là liên kết chúng theo một mô hình sự kiện có thể trao đổi và đối chiếu. GS1 EPCIS cung cấp cơ sở để tổ chức dữ liệu về những gì đã xảy ra với sản phẩm, thời điểm, địa điểm và tác nhân tham gia; đây là nền tảng phù hợp để đề tài xây dựng profile sự kiện truy xuất trong phạm vi PoC [5].")
add_para(doc, "Tuy nhiên, một lịch sử được ghi nhận chưa mặc nhiên là lịch sử nhất quán. Một sự kiện có thể hợp lệ về cấu trúc nhưng vẫn bất thường khi đối chiếu vị trí, chủ thể, trạng thái mã vùng trồng, sản lượng hoặc quan hệ chuyển giao với các dữ liệu và sự kiện liên quan. Ví dụ, mã vùng trồng có thể được sử dụng ngoài polygon đã đăng ký; sản lượng tích lũy có thể vượt năng lực đã khai báo; đầu ra của một công đoạn có thể không tương ứng với đầu vào có nguồn gốc hợp lệ; hoặc trình tự bàn giao có thể không phù hợp với chuỗi nghiệp vụ. Dữ liệu không gian, mô hình sự kiện và các cơ chế tăng hiệu quả truy xuất là những nền tảng quan trọng, song không tự xác thực tính hợp lý của nội dung provenance [5]-[7].")
add_para(doc, "Bên cạnh đó, tính bất biến của blockchain không giải quyết trực tiếp độ chính xác của dữ liệu hình thành ngoài chuỗi; một tọa độ, khối lượng hoặc bằng chứng không chính xác vẫn có thể được cam kết bất biến. Giới hạn này phù hợp với thảo luận về oracle paradox trong các hệ thống smart contract [10]. Do đó, blockchain trong đề tài không được xem là cơ chế chứng minh tuyệt đối tính đúng của dữ liệu đầu vào, mà là lớp bảo toàn bằng chứng cho dữ liệu đã đi qua các bước chuẩn hóa và kiểm tra.")
add_para(doc, "Từ khoảng trống giữa traceability và provenance integrity, đề tài BATS-AgriGuard đề xuất quy trình validation-before-anchoring. Mỗi vùng trồng được mô hình hóa bằng Digital Plantation Identity (DPI), liên kết mã vùng trồng với polygon địa lý, diện tích, loại cây trồng, chủ thể được ủy quyền, trạng thái và thời gian hiệu lực. Các sự kiện từ thu hoạch, tiếp nhận/chuyển giao, đóng gói đến xuất hàng được tổ chức theo hướng căn chỉnh với GS1 EPCIS. Plantation Code Integrity Engine (PCIE) kiểm tra các ràng buộc về không gian, quyền chủ thể, trạng thái, năng lực sản lượng, bằng chứng, Mass Balance và Chain of Custody ở cả cấp sự kiện lẫn xuyên sự kiện.")
add_para(doc, "Sau kiểm tra, dữ liệu được chuẩn hóa, băm, tổng hợp thành Merkle root và neo commitment lên blockchain EVM để hỗ trợ xác minh tính toàn vẹn sau thời điểm cam kết. Kết quả được phân loại thành PASS, REVIEW hoặc BLOCK cùng rule kích hoạt và dữ liệu liên quan, nhằm giúp người dùng truy xuất, giải thích và hậu kiểm. Các cảnh báo của hệ thống được định vị là dấu hiệu bất thường cần xem xét, không phải kết luận tự động về hành vi gian lận.")
add_figure(doc, "image1.png", "Hình 1. Kiến trúc tổng thể của hệ thống BATS-AgriGuard")
add_para(doc, "Dựa trên hướng tiếp cận trên, đề tài dự kiến đóng góp:", before=2, after=4, bold=True)
for item in [
    "Mô hình Digital Plantation Identity (DPI) biến mã số vùng trồng từ một chuỗi định danh đơn lẻ thành một thực thể số có ranh giới địa lý, diện tích, loại cây trồng, chủ thể được ủy quyền, trạng thái, thời gian hiệu lực và lịch sử sự kiện. Mô hình này tạo cơ sở để đối chiếu claim nguồn gốc với dữ liệu vùng trồng đã đăng ký.",
    "Plantation Code Integrity Engine (PCIE) với tám nhóm bất biến giải thích được: Geospatial Integrity, Identity/Authorization, Plantation Status, Seasonal Yield Capacity, Evidence Duplication, Spatio-temporal Consistency, Mass Balance và Chain of Custody. PCIE mở rộng kiểm tra từ tính hợp lệ của một sự kiện sang tính nhất quán giữa nhiều sự kiện, lô hàng và chủ thể trong chuỗi.",
    "Profile dữ liệu và luồng sự kiện căn chỉnh với GS1 EPCIS cho các nghiệp vụ trong phạm vi nghiên cứu, liên kết vùng trồng với sự kiện thu hoạch, tiếp nhận/chuyển giao, đóng gói và xuất hàng. Cách tổ chức này giúp tách rõ dữ liệu nghiệp vụ, dữ liệu không gian, bằng chứng và quan hệ custody để phục vụ trao đổi, truy vấn và hậu kiểm.",
    "Cơ chế quyết định PASS - REVIEW - BLOCK có thể giải thích: mỗi quyết định gắn với rule kích hoạt và dữ liệu liên quan. REVIEW được sử dụng để đưa các trường hợp cần xem xét thêm vào hàng đợi hậu kiểm, tránh đồng nhất dấu hiệu bất thường với kết luận gian lận.",
    "Kiến trúc validation-before-anchoring kết hợp lớp thu thập hiện trường, PostgreSQL/PostGIS, quản lý sự kiện và bằng chứng, cùng lớp Merkle/EVM. Dữ liệu chỉ được tạo commitment sau kiểm tra; blockchain vì vậy bảo toàn tính toàn vẹn sau cam kết và hỗ trợ inclusion verification, thay vì tự chứng minh tính đúng của dữ liệu đầu vào.",
    "Khung kiểm chứng cho PoC gồm các kịch bản hợp lệ, biên và đối kháng đối với từng nhóm bất biến; đồng thời tách việc đánh giá rule engine, truy vấn không gian và Merkle để làm rõ phạm vi, giới hạn và tính khả thi kỹ thuật của từng thành phần.",
]: add_bullet(doc, item)
add_heading(doc, "B2. Mục tiêu, nội dung, kế hoạch nghiên cứu", 2)
for item in [
    "Nghiên cứu mô hình hóa dữ liệu vùng trồng, sự kiện truy xuất và các quan hệ provenance theo hướng căn chỉnh với GS1 EPCIS.",
    "Thiết kế PCIE để kiểm tra tính toàn vẹn mã số vùng trồng và tính nhất quán dữ liệu ở cấp sự kiện lẫn xuyên sự kiện.",
    "Hiện thực PoC, kiểm chứng theo các kịch bản xác định và báo cáo rõ phạm vi, giới hạn, khả năng áp dụng của hệ thống.",
]: add_bullet(doc, item)
add_heading(doc, "B2.1. Mục tiêu", 3)
add_para(doc, "Mục tiêu tổng quát: Xây dựng và đánh giá BATS-AgriGuard, hệ thống truy xuất nguồn gốc nông sản dựa trên GS1 EPCIS, tập trung kiểm tra tính toàn vẹn của mã số vùng trồng và tính nhất quán nguồn gốc xuyên suốt chuỗi cung ứng.")
add_para(doc, "Mục tiêu cụ thể:", before=2, after=4, bold=True)
for item in [
    "Xây dựng mô hình DPI liên kết mã số vùng trồng với polygon địa lý, diện tích, loại cây trồng, chủ thể được ủy quyền, trạng thái và khoảng thời gian hiệu lực.",
    "Xây dựng và đánh giá PCIE với tám nhóm bất biến: Geospatial Integrity (G), Identity/Authorization (I), Plantation Status (S), Seasonal Yield Capacity (Y), Evidence Duplication (D), Spatio-temporal Consistency (T), Mass Balance (M) và Chain of Custody (C).",
    "Mô hình hóa các sự kiện truy xuất từ thu hoạch, chuyển giao, đóng gói đến xuất hàng theo hướng căn chỉnh với GS1 EPCIS, phục vụ trao đổi và truy vấn dữ liệu giữa các tác nhân.",
    "Xây dựng cơ chế PASS - REVIEW - BLOCK có giải thích, mở rộng kiểm tra từ từng sự kiện riêng lẻ sang tính nhất quán giữa nhiều sự kiện và nhiều chủ thể.",
    "Xây dựng cơ chế bảo toàn và xác minh bằng chứng bằng cryptographic hash, Merkle Tree và EVM blockchain anchoring; triển khai PoC và kiểm chứng trên kịch bản hợp lệ, biên và đối kháng.",
]: add_bullet(doc, item)

add_heading(doc, "B2.2. Nội dung và phương pháp nghiên cứu", 3)
add_heading(doc, "Nội dung 1: Khảo sát cơ sở lý thuyết và đặc tả mô hình dữ liệu", 3)
add_para(doc, "Mục tiêu: Xác định cơ sở khoa học, phạm vi hệ thống, tác nhân và lược đồ dữ liệu cần thiết để mô hình hóa mã số vùng trồng cùng lịch sử provenance.", after=4, bold=True)
add_para(doc, "Phương pháp thực hiện: Nghiên cứu tài liệu về truy xuất nguồn gốc nông sản, GS1 EPCIS/CBV, dữ liệu không gian, Mass Balance, Chain of Custody và blockchain; phân tích khoảng trống giữa traceability và provenance integrity. Trên cơ sở đó, đặc tả DPI gồm mã vùng trồng, polygon, diện tích, loại cây trồng, quyền sử dụng, trạng thái và thời gian hiệu lực; mô hình hóa actor, batch, sự kiện, bằng chứng và quan hệ chuyển giao.")
add_para(doc, "Kết quả dự kiến: Bộ đặc tả dữ liệu DPI và profile sự kiện căn chỉnh GS1 EPCIS cho chuỗi vùng trồng → thu hoạch → chuyển giao → đóng gói → xuất hàng; xác định rõ phạm vi PoC, các tác nhân và giới hạn nghiên cứu.")

add_heading(doc, "Nội dung 2: Xây dựng Plantation Code Integrity Engine (PCIE)", 3)
add_para(doc, "Mục tiêu: Xây dựng lớp kiểm tra có thể giải thích, mở rộng đánh giá từ từng sự kiện đơn lẻ sang mối quan hệ giữa nhiều sự kiện và nhiều chủ thể.", after=4, bold=True)
add_para(doc, "Phương pháp thực hiện: Xây dựng các luật kiểm tra giải thích được ở cấp sự kiện và xuyên sự kiện. Kiểm tra cấp sự kiện gồm geofence, quyền chủ thể, trạng thái/thời hạn mã vùng trồng, năng lực sản lượng theo mùa vụ và tái sử dụng bằng chứng theo policy. Kiểm tra xuyên sự kiện gồm tính nhất quán không gian - thời gian, Mass Balance và Chain of Custody. Kết quả được tổng hợp theo cơ chế PASS - REVIEW - BLOCK; risk score chỉ dùng để ưu tiên xem xét, không thay thế logic enforcement của các hard invariant.")
add_para(doc, "Kết quả dự kiến: PCIE trả về quyết định, rule kích hoạt và dữ liệu liên quan phục vụ giải thích; Review Queue lưu các trường hợp cần kiểm tra bổ sung. Mỗi cảnh báo được định vị là dấu hiệu bất thường cần hậu kiểm, không phải kết luận gian lận.")

add_heading(doc, "Nội dung 3: Hiện thực PoC truy xuất, kiểm tra và bảo toàn bằng chứng", 3)
add_para(doc, "Mục tiêu: Hiện thực một luồng vận hành tích hợp từ ghi nhận dữ liệu hiện trường đến truy xuất và xác minh bằng chứng sau kiểm tra.", after=4, bold=True)
add_para(doc, "Phương pháp thực hiện: Hiện thực lớp quản lý dữ liệu nghiệp vụ và không gian bằng PostgreSQL/PostGIS; lớp thu thập dữ liệu hiện trường trên Zalo Mini App; website quản trị, truy xuất và xác minh QR/GS1 Digital Link; backend điều phối sự kiện, validation result, evidence và audit log. Các sự kiện được kiểm tra trước khi hash, tổng hợp Merkle root và neo commitment lên mạng EVM. Bằng chứng được dùng cho inclusion verification, không thay cho cơ chế xác thực sự thật vật lý.")
add_para(doc, "Kết quả dự kiến: Một vertical slice có thể ghi nhận dữ liệu vùng trồng và sự kiện, chạy PCIE, lưu lịch sử truy xuất, xử lý REVIEW và công bố bằng chứng Merkle/EVM sau kiểm tra. EPCIS được định vị là EPCIS-aligned trong phạm vi các ObjectEvent hỗ trợ, không tuyên bố GS1 EPCIS conformance đầy đủ.")
add_heading(doc, "Nội dung 4: Kiểm chứng, đánh giá và hoàn thiện báo cáo", 3)
add_para(doc, "Mục tiêu: Kiểm chứng implementation theo các kịch bản xác định, đánh giá khả năng phát hiện các lớp vi phạm đã đặc tả và báo cáo rõ giới hạn của kết quả.", after=4, bold=True)
add_para(doc, "Phương pháp thực hiện: Thiết kế bộ kiểm thử độc lập gồm trường hợp hợp lệ, biên và đối kháng cho từng bất biến; dùng model-based testing cho Chain of Custody và property-based testing với seed cố định cho các tính chất của các rule. Kiểm tra profile EPCIS bằng JSON Schema chính thức trong phạm vi hỗ trợ; đánh giá metric TP, FP, TN, FN, Precision, Recall và F1 theo tập kịch bản. Thực hiện microbenchmark tách rời cho rule engine, truy vấn không gian PostGIS và Merkle; không diễn giải chúng thành latency end-to-end.")
add_para(doc, "Kết quả dự kiến: Báo cáo kiểm chứng implementation, coverage theo từng họ bất biến, phân tích ablation và độ nhạy các tham số sản lượng/Mass Balance; báo cáo benchmark riêng cho PCIE, PostGIS và Merkle, cùng thảo luận về hạn chế và hướng phát triển.")

doc.add_page_break()
add_heading(doc, "B2.3. Kế hoạch nghiên cứu", 3)
add_para(doc, "Bảng 1. Kế hoạch nghiên cứu", align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=4, bold=True, size=11)
add_grid_table(doc, ["Thời gian", "Nội dung chính", "Kết quả/đầu ra"], [
    ["Tháng 1", "Khảo sát tài liệu; xác định phạm vi, đối tượng, tác nhân và các khoảng trống nghiên cứu.", "Khung lý thuyết, phạm vi PoC và đặc tả yêu cầu."],
    ["Tháng 2", "Thiết kế DPI, dữ liệu vùng trồng, profile sự kiện và các bất biến PCIE.", "Đặc tả dữ liệu, rule và policy kiểm tra."],
    ["Tháng 3", "Hiện thực các kiểm tra cấp sự kiện; tích hợp PostgreSQL/PostGIS và quản lý evidence.", "PCIE phần kiểm tra G/I/S/Y/D; dữ liệu không gian và evidence."],
    ["Tháng 4", "Hiện thực kiểm tra xuyên sự kiện, Mass Balance, CoC và cơ chế PASS/REVIEW/BLOCK.", "PCIE phần T/M/C; Review Queue và evidence trail."],
    ["Tháng 5", "Tích hợp truy xuất, QR/GS1 Digital Link, Merkle commitment và EVM anchoring; xây dựng kịch bản kiểm thử.", "PoC tích hợp và bộ scenario kiểm chứng."],
    ["Tháng 6", "Đánh giá, phân tích kết quả, hoàn thiện báo cáo và sản phẩm trình bày.", "Báo cáo tổng kết, kết quả kiểm chứng và đề xuất hướng phát triển."],
], [1200, 5000, 3160], 10)

add_heading(doc, "B3. Kết quả dự kiến", 2)
for item in [
    "Mô hình DPI cho mã số vùng trồng, liên kết dữ liệu không gian, chủ thể, trạng thái, thời hạn hiệu lực và lịch sử sự kiện.",
    "PCIE thực thi tám nhóm bất biến G/I/S/Y/D/T/M/C ở cấp sự kiện và xuyên sự kiện, có giải thích cho kết quả PASS, REVIEW hoặc BLOCK.",
    "PoC BATS-AgriGuard gồm lớp thu thập dữ liệu hiện trường, website quản trị/truy xuất, backend, PostgreSQL/PostGIS, EPCIS Event Store, Review Queue, QR/GS1 Digital Link và lớp Merkle/EVM anchoring.",
    "Bộ kiểm thử mô tả rõ kịch bản hợp lệ, biên và đối kháng; báo cáo kiểm chứng implementation, metric phân loại, ablation, phân tích độ nhạy và microbenchmark theo từng thành phần.",
    "Báo cáo khoa học làm rõ đóng góp validation-before-anchoring, giới hạn của dữ liệu synthetic, giới hạn của blockchain đối với dữ liệu đầu vào, và các hạng mục cần hardening/pilot thực địa.",
]: add_bullet(doc, item)
add_para(doc, "Giới hạn cần nêu trong báo cáo: Hệ thống kiểm tra tính nhất quán của dữ liệu được cung cấp, không xác nhận chất lượng hoặc an toàn thực phẩm, không thay thế cơ quan quản lý/tổ chức chứng nhận/kiểm tra thực địa, và không coi bất thường là bằng chứng trực tiếp của hành vi gian lận.", before=6, after=8, italic=True)

add_heading(doc, "B4. Tài liệu tham khảo", 2)
refs = [
    "A. Ahmad and K. Bailey, “Blockchain in food traceability: A systematic literature review,” 2021 32nd Irish Signals and Systems Conference (ISSC), pp. 1-6, 2021.",
    "R. M. Ellahi, L. C. Wood, and A. E.-D. A. Bekhit, “Blockchain-based frameworks for food traceability: A systematic review,” Foods, vol. 12, no. 16, Art. no. 3026, 2023.",
    "A. Shahid et al., “Blockchain-based agri-food supply chain: A complete solution,” IEEE Access, vol. 8, pp. 69230-69243, 2020.",
    "M. P. Caro et al., “Blockchain-based traceability in Agri-Food supply chain management: A practical implementation,” IOT Tuscany, pp. 1-4, 2018.",
    "J. Ahn et al., “Oliot EPCIS: An open-source EPCIS 2.0 system for supply chain transparency,” SoftwareX, vol. 23, Art. no. 101477, 2023.",
    "L. Miloudi et al., “Smart sustainable farming management using integrated approach of IoT, blockchain & geospatial technologies,” International Conference on Advanced Intelligent Systems for Sustainable Development, pp. 340-347, 2019.",
    "H. Wu, S. Jiang, and J. Cao, “High-efficiency blockchain-based supply chain traceability,” IEEE Transactions on Intelligent Transportation Systems, vol. 24, no. 4, pp. 3748-3758, 2023.",
    "T. H. Pranto et al., “Blockchain and smart contract for IoT enabled smart agriculture,” PeerJ Computer Science, vol. 7, Art. no. e407, 2021.",
    "A. Spitalleri et al., “BioTrak: A blockchain-based platform for food chain logistics traceability,” ICCNS, pp. 105-110, 2023.",
    "A. Albizri and D. Appelbaum, “Trust but verify: The oracle paradox of blockchain smart contracts,” Journal of Information Systems, vol. 35, no. 2, pp. 1-16, 2021.",
]
for i, ref in enumerate(refs, 1): add_reference(doc, i, ref)

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
sign = doc.add_table(rows=1, cols=2)
no_cell_borders(sign)
fill_cell(sign.cell(0, 0), "Ngày ..... tháng ..... năm 2026\nGIẢNG VIÊN HƯỚNG DẪN\n(Ký và ghi rõ họ tên)\n\n\n\n", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
fill_cell(sign.cell(0, 1), "Ngày ..... tháng ..... năm 2026\nCHỦ NHIỆM ĐỀ TÀI\n(Ký và ghi rõ họ tên)\n\n\n\nNguyễn Đình Khang", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
set_table_geometry(sign, [4680, 4680])

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run()
    set_font(run, size=10)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)

doc.core_properties.title = "BATS-AgriGuard - Dự thảo thuyết minh đề tài NCKH sinh viên 2026"
doc.core_properties.subject = "Chuyển hóa từ nội dung công trình dự thi Eureka 2026"
doc.core_properties.author = "Nguyễn Đình Khang; Võ Minh An"
doc.save(OUT)
print(OUT.resolve())
