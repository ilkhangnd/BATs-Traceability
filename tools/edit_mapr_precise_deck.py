from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SRC = Path("/private/tmp/codex-presentations/manual-mapr-precise-edit/tmp/source.pptx")
OUT = Path("/Users/nguyendinhkhang/khangnd/BATs/outputs/MAPR2026-PRECISE-conference-polished.pptx")

BLUE = RGBColor(0, 83, 166)
DEEP_BLUE = RGBColor(0, 62, 125)
LIGHT_BLUE = RGBColor(226, 240, 255)
PALE_BLUE = RGBColor(242, 247, 253)
GREEN = RGBColor(0, 154, 112)
DARK = RGBColor(29, 58, 96)
GRAY = RGBColor(90, 106, 130)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(205, 70, 70)


def set_text(shape, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, font="Agrandir"):
    shape.text = text
    tf = shape.text_frame
    tf.clear()
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.font.name = font
    return shape


def add_text(slide, x, y, w, h, text, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, font="Agrandir"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.margin_left = Inches(0.08)
    shape.text_frame.margin_right = Inches(0.08)
    shape.text_frame.margin_top = Inches(0.04)
    shape.text_frame.margin_bottom = Inches(0.04)
    set_text(shape, text, size=size, color=color, bold=bold, align=align, font=font)
    return shape


def add_card(slide, x, y, w, h, fill=PALE_BLUE, line=LIGHT_BLUE, radius=True):
    geom = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(geom, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.1)
    return shape


def add_chip(slide, x, y, w, text, icon=None, fill=WHITE, text_color=BLUE):
    add_card(slide, x, y, w, 0.42, fill=fill, line=LIGHT_BLUE)
    label = f"{icon}  {text}" if icon else text
    return add_text(slide, x + 0.05, y + 0.055, w - 0.1, 0.32, label, size=11, color=text_color, bold=True, align=PP_ALIGN.CENTER)


def add_circle_icon(slide, x, y, label, diameter=0.42, fill=BLUE, font_size=9):
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    icon.fill.solid()
    icon.fill.fore_color.rgb = fill
    icon.line.color.rgb = fill
    tx = add_text(slide, x, y + 0.055, diameter, diameter - 0.08, label, size=font_size, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    tx.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return icon


def add_badge_label(slide, x, y, badge, label, w=4.2, size=15, fill=BLUE):
    add_circle_icon(slide, x, y, badge, fill=fill, font_size=8 if len(badge) > 3 else 9)
    return add_text(slide, x + 0.5, y + 0.06, w, 0.34, label, size=size, color=DARK, bold=True)


def add_key_finding(slide, text, x=3.0, y=9.05, w=14.0):
    add_card(slide, x, y, w, 0.72, fill=RGBColor(237, 250, 245), line=RGBColor(180, 226, 209))
    add_text(slide, x + 0.25, y + 0.1, 2.0, 0.34, "Key Finding", size=12, color=GREEN, bold=True)
    add_text(slide, x + 2.18, y + 0.08, w - 2.45, 0.46, f"✓ {text}", size=14, color=DARK, bold=True)


def clear_shape_text(shape):
    if hasattr(shape, "text_frame"):
        shape.text_frame.clear()


def style_existing_text(shape, font_size=None, color=None, bold=None, font="Agrandir"):
    if not hasattr(shape, "text_frame"):
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if font_size:
                r.font.size = Pt(font_size)
            if color:
                r.font.color.rgb = color
            if bold is not None:
                r.font.bold = bold
            r.font.name = font


def slide3(slide):
    clear_shape_text(slide.shapes[6])
    rows = [
        ("ERC", "ERC-20 assumptions shape AMM behavior."),
        ("DeFi", "DeFi protocols rely on balance invariants."),
        ("MEV", "MEV emerges when token state and price diverge."),
    ]
    y = 3.05
    for badge, line in rows:
        add_badge_label(slide, 4.1, y, badge, line, w=10.7, size=14)
        y += 0.52

    clear_shape_text(slide.shapes[13])
    add_card(slide, 1.25, 6.85, 8.0, 1.85, fill=LIGHT_BLUE, line=RGBColor(137, 185, 233))
    add_text(slide, 1.55, 6.98, 7.4, 1.55, "Token state changes\n↓\nAMM price remains unchanged\n↓\nPrice inconsistency", size=13, color=DEEP_BLUE, bold=True, align=PP_ALIGN.CENTER)

    add_chip(slide, 16.25, 8.63, 2.2, "Searcher Profit", icon="$", fill=RGBColor(237, 250, 245), text_color=GREEN)


def slide4(slide):
    add_circle_icon(slide, 0.92, 2.20, "SA", fill=BLUE, font_size=9)
    add_circle_icon(slide, 10.55, 2.20, "DF", fill=BLUE, font_size=9)
    add_circle_icon(slide, 0.92, 7.48, "RO", fill=GREEN, font_size=9)
    slide.shapes[7].left = Inches(1.55)
    slide.shapes[7].width = Inches(7.1)
    slide.shapes[9].left = Inches(10.95)
    slide.shapes[9].width = Inches(7.0)
    slide.shapes[12].left = Inches(1.55)
    slide.shapes[12].width = Inches(16.9)


def slide5(slide):
    clear_shape_text(slide.shapes[9])
    add_card(slide, 9.35, 5.95, 8.45, 3.0, fill=RGBColor(246, 251, 255), line=RGBColor(183, 211, 240))
    add_text(slide, 9.72, 6.18, 7.8, 0.38, "Research Gap", size=16, color=DEEP_BLUE, bold=True)
    gap_items = [("×", "Static only", RED), ("×", "Dynamic only", RED), ("×", "No replayable witness", RED), ("✓", "PRECISE integrates all.", GREEN)]
    y = 6.72
    for mark, label, color in gap_items:
        add_text(slide, 9.85, y, 0.35, 0.3, mark, size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, 10.30, y, 7.1, 0.3, label, size=13.5, color=DARK, bold=(mark == "✓"))
        y += 0.42


def slide6(slide):
    # Pull the large architecture figure up slightly to make room for the memory hook.
    fig = slide.shapes[8]
    fig.top = Inches(1.95)
    fig.height = Inches(5.55)
    slide.shapes[7].top = Inches(7.60)

    add_card(slide, 1.25, 8.18, 17.5, 1.22, fill=RGBColor(246, 251, 255), line=RGBColor(188, 215, 243))
    add_text(slide, 1.55, 8.30, 2.1, 0.34, "Key Features", size=13, color=DEEP_BLUE, bold=True)
    features = [
        ("SDG", "Static SDG Pruning"),
        ("FORK", "Fork-based Fuzzing"),
        ("DIFF", "State Difference Verification"),
        ("$", "Economic Validation"),
    ]
    x = 3.55
    widths = [3.45, 3.45, 4.25, 3.45]
    for (icon, label), w in zip(features, widths):
        add_chip(slide, x, 8.70, w, label, icon=icon, fill=WHITE, text_color=BLUE)
        x += w + 0.27


def slide7(slide):
    # Outline rings keep the original text intact while calling attention to the sampling weights.
    for y in [4.42, 4.85, 5.28]:
        ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.65), Inches(y), Inches(0.82), Inches(0.36))
        ring.fill.background()
        ring.line.color.rgb = GREEN
        ring.line.width = Pt(1.5)


def slide8(slide):
    add_circle_icon(slide, 0.82, 1.63, "T1", fill=RED, font_size=9)
    add_circle_icon(slide, 0.82, 5.38, "T2", fill=BLUE, font_size=9)
    slide.shapes[5].left = Inches(1.55)
    slide.shapes[7].left = Inches(1.55)
    formula = slide.shapes[9]
    formula.left = Inches(6.55)
    formula.top = Inches(8.95)
    formula.width = Inches(6.75)
    formula.height = Inches(0.70)


def slide14(slide):
    chart = slide.shapes[4]
    chart.left = Inches(1.05)
    chart.top = Inches(2.00)
    chart.width = Inches(10.8)
    chart.height = Inches(6.15)
    slide.shapes[8].left = Inches(1.05)
    slide.shapes[8].top = Inches(8.35)
    slide.shapes[8].width = Inches(10.8)

    add_card(slide, 12.25, 2.05, 5.95, 6.75, fill=RGBColor(246, 251, 255), line=RGBColor(188, 215, 243))
    add_text(slide, 12.62, 2.35, 5.25, 0.48, "Top Root Causes", size=18, color=DEEP_BLUE, bold=True)
    causes = [("SDG", "SDG slicing"), ("ADM", "Admin adjustment"), ("PRX", "Dynamic proxy"), ("POOL", "Pool context")]
    y = 3.28
    for icon, label in causes:
        add_card(slide, 12.65, y, 5.15, 0.82, fill=WHITE, line=LIGHT_BLUE)
        add_circle_icon(slide, 12.88, y + 0.18, icon, diameter=0.42, fill=BLUE, font_size=7 if len(icon) > 3 else 8)
        add_text(slide, 13.48, y + 0.18, 4.0, 0.36, label, size=15, color=DARK, bold=True)
        y += 1.08


def slide16(slide):
    clear_shape_text(slide.shapes[7])
    add_card(slide, 1.18, 2.02, 8.45, 4.62, fill=RGBColor(246, 251, 255), line=RGBColor(184, 213, 241))
    add_circle_icon(slide, 1.48, 2.33, "S", fill=BLUE, font_size=9)
    add_text(slide, 2.00, 2.30, 7.25, 0.50, "Scope", size=18, color=DEEP_BLUE, bold=True)
    scope = ["Ethereum Mainnet only", "D = 3", "Max 3-call sequences"]
    y = 3.16
    for item in scope:
        add_text(slide, 1.72, y, 0.35, 0.32, "•", size=16, color=GREEN, bold=True)
        add_text(slide, 2.15, y, 6.8, 0.34, item, size=15.5, color=DARK, bold=True)
        y += 0.72

    add_card(slide, 10.35, 2.02, 8.45, 4.62, fill=RGBColor(246, 251, 255), line=RGBColor(184, 213, 241))
    add_circle_icon(slide, 10.65, 2.33, "!", fill=RED, font_size=11)
    add_text(slide, 11.15, 2.30, 7.25, 0.50, "Practical Constraints", size=18, color=DEEP_BLUE, bold=True)
    constraints = ["Constant-product AMMs", "No builder competition", "No private order flow"]
    y = 3.16
    for item in constraints:
        add_text(slide, 10.88, y, 0.35, 0.32, "•", size=16, color=GREEN, bold=True)
        add_text(slide, 11.32, y, 6.8, 0.34, item, size=15.5, color=DARK, bold=True)
        y += 0.72

    add_card(slide, 5.15, 7.25, 9.85, 1.55, fill=RGBColor(237, 250, 245), line=RGBColor(178, 226, 207))
    add_circle_icon(slide, 5.52, 7.56, "!", fill=GREEN, font_size=11)
    add_text(slide, 6.15, 7.35, 8.4, 0.90, "Discoverability\ndoes not imply\nReal-world Profitability", size=16, color=DEEP_BLUE, bold=True, align=PP_ALIGN.CENTER)


def slide17(slide):
    clear_shape_text(slide.shapes[7])
    steps = [
        ("H", "Hybrid Detection", "SDG + Fork Fuzzing"),
        ("S", "State Verification", ""),
        ("P", "Profitability Validation", ""),
        ("T", "Scalable tMEV Discovery", ""),
    ]
    x, y = 1.35, 1.92
    for i, (icon, title, sub) in enumerate(steps):
        add_card(slide, x, y, 8.2, 1.05, fill=RGBColor(246, 251, 255), line=RGBColor(184, 213, 241))
        add_circle_icon(slide, x + 0.28, y + 0.24, icon, fill=BLUE, font_size=9)
        add_text(slide, x + 0.90, y + 0.16, 6.8, 0.30, title, size=15.5, color=DEEP_BLUE, bold=True)
        if sub:
            add_text(slide, x + 0.90, y + 0.55, 6.8, 0.26, sub, size=12.8, color=GRAY)
        if i < len(steps) - 1:
            add_text(slide, x + 3.85, y + 0.98, 0.5, 0.34, "↓", size=17, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        y += 1.38

    add_card(slide, 10.25, 1.92, 7.95, 4.85, fill=RGBColor(237, 250, 245), line=RGBColor(178, 226, 207))
    add_text(slide, 10.60, 2.22, 7.2, 0.42, "Key Achievement", size=17, color=GREEN, bold=True)
    metrics = ["22,279 contracts", "84.58% Recall", "78.04% F1", "Ecosystem-scale deployment"]
    y = 3.00
    for metric in metrics:
        add_text(slide, 10.70, y, 0.38, 0.30, "✓", size=15, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, 11.15, y, 6.55, 0.32, metric, size=14.8, color=DARK, bold=True)
        y += 0.62

    add_card(slide, 1.35, 7.42, 16.85, 1.28, fill=LIGHT_BLUE, line=RGBColor(137, 185, 233))
    add_text(slide, 1.85, 7.72, 15.85, 0.55, "Static pruning + execution-grounded validation is an effective solution for scalable Token-centric MEV detection.", size=15.5, color=DEEP_BLUE, bold=True, align=PP_ALIGN.CENTER)


def slide18(slide):
    add_text(slide, 1.25, 6.90, 6.35, 0.68, "Questions?", size=28, color=GREEN, bold=True, align=PP_ALIGN.LEFT)


def main():
    prs = Presentation(str(SRC))
    slide3(prs.slides[2])
    slide4(prs.slides[3])
    slide5(prs.slides[4])
    slide6(prs.slides[5])
    slide7(prs.slides[6])
    slide8(prs.slides[7])
    add_key_finding(prs.slides[9], "Highest recall across both datasets.", y=9.10)
    add_key_finding(prs.slides[10], "Highest recall with competitive execution time.", y=9.15)
    add_key_finding(prs.slides[11], "SDG pruning provides the greatest runtime reduction.", y=8.62)
    add_key_finding(prs.slides[12], "D = 3 and N = 1000 offer the best trade-off.", y=9.10)
    slide14(prs.slides[13])
    slide16(prs.slides[15])
    slide17(prs.slides[16])
    slide18(prs.slides[17])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
